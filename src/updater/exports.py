"""Streaming current and history workbooks from one lifetime-pinned snapshot."""
from datetime import datetime
from collections.abc import Iterable, Sequence
import json
import os
from pathlib import Path
import tempfile
from typing import TypedDict
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from src.updater.reader import PublishedSnapshot
from src.updater.store import StoreError
from src.updater.source import ItemSource, observations_from_page
from core.catalog_sources import SearchQuery
from core.updater.models import ItemObservation


def export_xlsx(snapshot: PublishedSnapshot, destination: Path, *, sheet_rows: int = 1048576) -> None:
    if sheet_rows < 2 or sheet_rows > 1048576:
        raise StoreError('invalid_sheet_limit')
    book = Workbook(write_only=True)
    current_fields = ('listing_key','goods_id','display_name','seller_name','manufacturer_name','spec_raw',
              'contract_price_decimal','unit','currency','presence_status','item_version_id','first_seen_at',
              'last_checked_at','data_changed_at','price_changed_at','published_at','information_groups',
              'catalog_generation','profile_hash','snapshot_id','contract_number','contract_start','contract_end')
    changes = snapshot.connection.execute('SELECT * FROM item_changes ORDER BY observed_at,listing_key,field_path')
    prices = snapshot.connection.execute('SELECT * FROM price_history ORDER BY observed_at,listing_key,event_id')
    datasets = (('상품',current_fields,snapshot.current()),
                ('필드변경',tuple(column[0] for column in changes.description),(dict(row) for row in changes)),
                ('가격이력',tuple(column[0] for column in prices.description),(dict(row) for row in prices)))
    for name, fields, rows in datasets:
        sheet = book.create_sheet(name + '1')
        sheet.append(fields)
        row_number, index = 1, 1
        for item in rows:
            if row_number >= sheet_rows:
                index += 1
                sheet = book.create_sheet(name + str(index))
                sheet.append(fields)
                row_number = 1
            cells = []
            for field in fields:
                value = item.get(field)
                if value is not None:
                    if field.endswith('_at'):
                        value = datetime.fromisoformat(str(value)).astimezone(ZoneInfo('Asia/Seoul')).isoformat() + ' [Asia/Seoul]'
                    if isinstance(value, dict):
                        value = json.dumps(value, ensure_ascii=False)
                    value = str(value)
                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = 's'
                cells.append(cell)
            sheet.append(cells)
            row_number += 1
    book.save(destination)


class SampleReport(TypedDict):
    mode: str
    coverage_complete: bool
    source_total: int
    requested_limit: int
    collected_items: int
    out: str
    file_bytes: int


def export_sample(source: ItemSource, destination: Path, *, limit: int = 100) -> SampleReport:
    if not 1 <= limit <= 1000:
        raise StoreError('invalid_sample_limit')
    if destination.suffix.lower() != '.xlsx':
        raise StoreError('unsupported_sample_format')
    if destination.exists() or destination.is_symlink():
        raise StoreError('output_exists')
    items: list[ItemObservation] = []
    seen: set[str] = set()
    total = None
    size = min(100, limit)
    for number in range(1, (limit+size-1)//size+1):
        page = source.fetch(SearchQuery(page=number, page_size=size))
        total = page.total_size if total is None else total
        expected = min(size, max(0, total-(number-1)*size))
        observations = observations_from_page(page)
        if page.total_size != total or page.row_total not in (None, total) or len(observations) != expected:
            raise StoreError('sample_count_drift_or_short_page')
        for item in observations:
            if item.listing_key in seen:
                raise StoreError('sample_duplicate_item')
            seen.add(item.listing_key)
        items.extend(observations[:limit-len(items)])
        if len(items) >= min(limit, total):
            break
    report = SampleReport(mode='partial_test', coverage_complete=False,
                               source_total=total or 0, requested_limit=limit,
                               collected_items=len(items), out=str(destination), file_bytes=0)
    write_observation_xlsx(items, ((key,str(value)) for key,value in report.items() if key!='file_bytes'),
                           destination, metadata_sheet='테스트정보', items_sheet='상품_테스트')
    report['file_bytes'] = destination.stat().st_size
    return report


def write_observation_xlsx(items: Sequence[ItemObservation], metadata: Iterable[tuple[str,str]],
                           destination: Path, *, metadata_sheet: str, items_sheet: str) -> None:
    book = Workbook()
    book.remove(book.worksheets[0])
    info = book.create_sheet(metadata_sheet)
    info.append(['field', 'value'])
    for pair in metadata:
        cells = []
        for text in pair:
            cell = WriteOnlyCell(info, value=text)
            cell.data_type = 's'
            cells.append(cell)
        info.append(cells)
    sheet = book.create_sheet(items_sheet)
    sheet.append(['listing_key', 'goods_id', 'display_name', 'seller_name', 'manufacturer_name',
                  'spec_raw', 'contract_price_decimal', 'unit', 'currency', 'observed_at'])
    for item in items:
        cells = []
        for value in (item.listing_key, item.goods_id, item.display_name, item.seller_name,
                      item.manufacturer_name, item.spec_raw, item.contract_price_decimal,
                      item.unit, item.currency, item.observed_at.astimezone(ZoneInfo('Asia/Seoul')).isoformat()):
            cell = WriteOnlyCell(sheet, value=None if value is None else str(value))
            cell.data_type = 's'
            cells.append(cell)
        sheet.append(cells)
    descriptor, temporary = tempfile.mkstemp(prefix=destination.name+'.', suffix='.partial', dir=destination.parent)
    os.close(descriptor)
    try:
        book.save(temporary)
        os.link(temporary, destination)
    finally:
        book.close()
        Path(temporary).unlink(missing_ok=True)
