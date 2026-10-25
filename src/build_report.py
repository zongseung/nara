import io
import os
import runpy
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.comparison import compare
from src.catalog_provider import LiveProvider

ROWS = ["물품식별번호", "규 격", "모 델 명", "재 질", "용 도", "가 격", "사 진"]

HEADER_FILL = PatternFill("solid", fgColor="D9D9F3")
OUR_FILL = PatternFill("solid", fgColor="FFFF00")
APPROX_FILL = PatternFill("solid", fgColor="FCE4D6")   # 근사(band 폴백) 후보 강조: 연주황
THIN = Side(style="thin", color="9AA0B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)


def _fmt(our, key):
    """우리 품목 dict → 구분 라벨에 해당하는 표시값."""
    if key == "물품식별번호":
        return our.get("식별번호", "")
    if key == "규 격":
        d = our.get("규격")
        return our.get("규격원문") or (f"{d[0]}×{d[1]}mm" if d else "원천 미기재")
    if key == "모 델 명":
        return our.get("모델", "")
    if key == "재 질":
        return our.get("재질원문") or our.get("재질") or "원천 미기재"
    if key == "용 도":
        faces = "양면" if our.get("표시면수") == 2 else "단면" if our.get("표시면수") == 1 else ""
        supports = our.get("지지대수")
        parts = [our.get("용도") or "원천 미기재", our.get("속성", {}).get("옵션/기타") or faces,
                 f"지지대 {supports}개" if supports is not None else "지지대 개수 미확인"]
        return " / ".join(part for part in parts if part)
    if key == "가 격":
        amount = our.get("가격")
        return f"{amount:,}" if amount is not None else "원천 미기재"
    return ""


def _put(ws, r, col, val, fill=None, bold=False):
    if isinstance(val, str) and val.startswith(("=", "+", "-", "@")):
        val = "'" + val
    cell = ws.cell(row=r, column=col, value=val)
    cell.alignment = CENTER
    cell.border = BORDER
    if fill:
        cell.fill = fill
    if bold:
        cell.font = BOLD
    return cell


def _add_image(ws, r, col, url, session):
    if not url:
        _put(ws, r, col, "원천 이미지 미기재")
        return
    try:
        full = url if url.startswith("http") else "https://shop.g2b.go.kr" + url
        with session.get(full, timeout=30) as response:
            response.raise_for_status()
            img = XLImage(io.BytesIO(response.content))
        scale = min(180 / img.width, 160 / img.height, 1)
        img.width, img.height = round(img.width * scale), round(img.height * scale)
        ws.add_image(img, f"{get_column_letter(col)}{r}")
    except (requests.RequestException, OSError, ValueError):
        _put(ws, r, col, "(이미지 실패)")


def prepare_output(out_path: str | Path) -> Path:
    path = Path(out_path).expanduser()
    if path.suffix.lower() != '.xlsx':
        raise ValueError(f"출력 파일의 확장자는 .xlsx여야 합니다: {path}")
    if path.exists() or path.is_symlink():
        raise FileExistsError(f"출력 경로가 이미 있습니다. 새 파일명을 지정하세요: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def build(our_items, out_path, top_n=3, our_company="제안업체", exclude_companies=None,
          exact_price=False, fallback_band=0.0, provider=None):
    output = prepare_output(out_path)
    excl = exclude_companies if exclude_companies is not None else [our_company]
    source = provider or LiveProvider()
    pairs = [compare(it, source, top_n=top_n, exact_price=exact_price, fallback_band=fallback_band, excluded=excl)
             for it in our_items]
    wb = Workbook(); ws = wb.worksheets[0]; ws.title = "가격비교표"
    with requests.Session() as session:
        _render_pairs(ws, our_company, pairs, top_n, session, provider)
        _save(wb, output)
    return out_path


def build_multi(companies, out_path, top_n=3, exact_price=True,
                exclude_companies=None, fallback_band=0.0, workers=8, provider=None):
    """{업체명: [our_item...]} → 업체별 시트. 동시(ThreadPool) 매칭 후 렌더. 기본 계열사 상호 제외.
    provider=None이면 실시간 쇼핑몰을 조회하고 사진을 내려받아 삽입한다."""
    output = prepare_output(out_path)
    excl = exclude_companies if exclude_companies is not None else list(companies)
    source = provider or LiveProvider()
    tasks = [(name, i, it) for name, items in companies.items() for i, it in enumerate(items)]

    def work(t):
        name, i, it = t
        it, cands = compare(it, source, top_n=top_n, exact_price=exact_price,
                            fallback_band=fallback_band, excluded=excl)
        mark = " (근사)" if cands and cands[0].get("근사") else ""
        print(f"[{name}] {it.get('모델') or it.get('품명')} ({it.get('가격') or 0:,}) → 후보 {len(cands)}개{mark}", flush=True)
        if it.get("비교보류사유"):
            print(f"[{name}] {it['비교보류사유']}", flush=True)
        return name, i, it, cands

    results = {}
    with ThreadPoolExecutor(max_workers=workers) as ex:   # 품목 단위 동시 매칭
        for name, i, it, cands in ex.map(work, tasks):
            results.setdefault(name, {})[i] = (it, cands)

    wb = Workbook(); wb.remove(wb.worksheets[0])
    with requests.Session() as session:
        for name, items in companies.items():
            ws = wb.create_sheet(name[:31])
            _render_pairs(ws, name, [results[name][i] for i in range(len(items))], top_n, session, provider)
        _save(wb, output)
    return out_path


def _save(workbook, out_path):
    path = Path(out_path)
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx") as staged:
            workbook.save(staged.name)
            os.link(staged.name, path)
    finally:
        workbook.close()


def _render_pairs(ws, our_company, pairs, top_n, session, provider=None):
    ncol = 3 + 1 + top_n                      # 연번/품명/구분 + 우리 + 후보N

    # 제목/업체명
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=ncol)
    t = _put(ws, 1, 1, "가격 비교표", bold=True)
    t.font = Font(bold=True, size=16)
    _put(ws, 2, 1, f"○ 업체명 : {our_company}")

    r = 3
    if provider is not None:
        _put(ws, r, 1, provider.description())
        r += 1
    for i, (our, cands) in enumerate(pairs, 1):
        cands = sorted(cands, key=lambda c: c.get("가격") or 0)
        if provider is not None and cands:
            _put(ws, r, 1, str(cands[0].get("catalog_metadata", {})))
            r += 1
        # 블록 헤더 (연번/품명/구분/회사명들). band 폴백 후보는 (근사) 표기.
        _put(ws, r, 1, "연번", HEADER_FILL, bold=True)
        _put(ws, r, 2, "품 명", HEADER_FILL, bold=True)
        _put(ws, r, 3, "구분", HEADER_FILL, bold=True)
        _put(ws, r, 4, our_company, OUR_FILL, bold=True)
        for j, c in enumerate(cands):
            ap = c.get("근사")
            header = _put(ws, r, 5 + j, c["업체"] + (" (근사)" if ap else ""),
                          APPROX_FILL if ap else HEADER_FILL, bold=True)
            if c.get("매칭근거"):
                header.comment = Comment(c["매칭근거"], "비교 근거")
        for j in range(len(cands), top_n):
            _put(ws, r, 5 + j, our.get("비교보류사유") or "조건 충족 후보 없음", HEADER_FILL)

        # 상세 7행
        for k, key in enumerate(ROWS):
            rr = r + 1 + k
            _put(ws, rr, 3, key)                              # 구분 라벨
            if key == "사 진":
                ws.row_dimensions[rr].height = 125
                _put(ws, rr, 4, "")
                if provider is None:
                    _add_image(ws, rr, 4, our.get("imgSrc"), session)
                else:
                    _put(ws, rr, 4, our.get("imgSrc") or "")
                for j, c in enumerate(cands):
                    _put(ws, rr, 5 + j, "")
                    if provider is None:
                        _add_image(ws, rr, 5 + j, c.get("imgSrc"), session)
                    else:
                        _put(ws, rr, 5 + j, c.get("imgSrc") or "")
            else:
                _put(ws, rr, 4, _fmt(our, key), bold=(key == "가 격"))
                for j, c in enumerate(cands):
                    cell = _put(ws, rr, 5 + j, _fmt(c, key), bold=(key == "가 격"))
                    if key == "가 격" and c.get("근사"):     # 근사 후보 단가 강조(우리와 다른 단가)
                        cell.fill = APPROX_FILL
                for j in range(len(cands), top_n):
                    _put(ws, rr, 5 + j, "")

        # 연번/품명 세로 병합 (상세 7행)
        ws.merge_cells(start_row=r + 1, end_row=r + 7, start_column=1, end_column=1)
        ws.merge_cells(start_row=r + 1, end_row=r + 7, start_column=2, end_column=2)
        _put(ws, r + 1, 1, i, bold=True)
        _put(ws, r + 1, 2, our["품명"], bold=True)

        r += 8

    # 열 너비
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    for col in range(4, ncol + 1):
        ws.column_dimensions[get_column_letter(col)].width = 28


def main():
    runpy.run_module("src.compare_file", run_name="__main__")


if __name__ == "__main__":
    main()
