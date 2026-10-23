"""Streaming diagnostic/legacy export with provenance and atomic destination replacement."""

import argparse
import json
import os
import sqlite3
import tempfile
from collections.abc import Iterator
from contextlib import closing
from dataclasses import asdict
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook

from src.catalog_provider import SQLiteProvider
from core.catalog_sources import SourceError
from core.catalog_json import json_text

BATCH = 1000
XLSX_ROWS = 1_048_576
_COLUMNS = (
    "run_id",
    "scope",
    "profile_hash",
    "entity_key",
    "entity_kind",
    "source_id",
    "goods_id",
    "observed_at",
    "raw_json",
    "validation",
)


def export_catalog(
    db: str | Path,
    out: str | Path,
    *,
    scope: str = "shopping-all",
    run: str = "last-reconciled",
    allow_incomplete: bool = False,
    xlsx_rows: int = XLSX_ROWS,
) -> int:
    """Preserve all raw keys in JSON; legacy normalized data is explicitly unverified."""
    if not 2 <= xlsx_rows <= XLSX_ROWS:
        raise SourceError("invalid_xlsx_row_limit")
    destination = Path(out)
    if destination.suffix.lower() not in (".xlsx", ".parquet"):
        raise SourceError("unsupported_export_format")
    if destination.resolve() == Path(db).resolve():
        raise SourceError("export_would_overwrite_database")
    with closing(
        sqlite3.connect(Path(db).resolve().as_uri() + "?mode=ro", uri=True)
    ) as con:
        con.row_factory = sqlite3.Row
        legacy = con.execute("PRAGMA user_version").fetchone()[0] != 2
        if legacy:
            if not con.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='items'"
            ).fetchone():
                raise SourceError("incompatible_database")
            metadata = {
                "validation": "legacy_unverified",
                "source_snapshot_verified": False,
                "raw_original_available": False,
            }
            cursor = con.execute("SELECT id,data FROM items ORDER BY id")

            def rows() -> Iterator[tuple[str | None, ...]]:
                for row in cursor:
                    yield (
                        None,
                        "legacy",
                        None,
                        row["id"],
                        None,
                        None,
                        None,
                        None,
                        row["data"],
                        "legacy_unverified",
                    )
        else:
            provider = SQLiteProvider(db, run, scope, allow_incomplete)
            metadata = asdict(provider.metadata)
            cursor = con.execute(
                "SELECT e.entity_key,e.goods_id,o.* FROM run_entities e JOIN observations o USING(observation_id) WHERE e.run_id=? ORDER BY e.entity_key",
                (provider.metadata.run_id,),
            )

            def rows() -> Iterator[tuple[str | None, ...]]:
                for row in cursor:
                    yield (
                        row["run_id"],
                        scope,
                        provider.metadata.profile_hash,
                        row["entity_key"],
                        row["entity_kind"],
                        row["source_id"],
                        row["goods_id"],
                        row["observed_at"],
                        row["raw_json"],
                        provider.metadata.validation,
                    )

        fd, temporary = tempfile.mkstemp(
            prefix=destination.name + ".", suffix=".partial", dir=destination.parent
        )
        os.close(fd)
        try:
            count = 0
            if destination.suffix.lower() == ".parquet":
                schema = pa.schema(
                    [(name, pa.string()) for name in _COLUMNS],
                    metadata={b"catalog": json_text(metadata).encode()},
                )
                with pq.ParquetWriter(temporary, schema, compression="zstd") as writer:
                    batch = []
                    for row in rows():
                        batch.append(dict(zip(_COLUMNS, row)))
                        count += 1
                        if len(batch) >= BATCH:
                            writer.write_table(
                                pa.Table.from_pylist(batch, schema=schema)
                            )
                            batch.clear()
                    if batch:
                        writer.write_table(pa.Table.from_pylist(batch, schema=schema))
            else:
                wb = Workbook(write_only=True)
                info = wb.create_sheet("provenance")
                info.append(["field", "value"])
                for key, value in metadata.items():
                    info.append([key, str(value)])
                sheet, used = None, xlsx_rows
                for row in rows():
                    if used >= xlsx_rows:
                        sheet = wb.create_sheet("data_" + str(len(wb.worksheets)))
                        sheet.append(_COLUMNS)
                        used = 1
                    assert sheet is not None
                    sheet.append(
                        [
                            "'" + v
                            if isinstance(v, str) and v.startswith(("=", "+", "-", "@"))
                            else v
                            for v in row
                        ]
                    )
                    count += 1
                    used += 1
                if sheet is None:
                    wb.create_sheet("data_1").append(_COLUMNS)
                wb.save(temporary)
                wb.close()
            with open(temporary, "rb") as saved:
                os.fsync(saved.fileno())
            os.replace(temporary, destination)
            return count
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", default="catalog_v2.db")
    parser.add_argument("--out", default="catalog.parquet")
    parser.add_argument("--scope", default="shopping-all")
    parser.add_argument("--run", default="last-reconciled")
    parser.add_argument("--allow-incomplete", action="store_true")
    args = parser.parse_args()
    try:
        count = export_catalog(
            args.db,
            args.out,
            scope=args.scope,
            run=args.run,
            allow_incomplete=args.allow_incomplete,
        )
        print(json.dumps({"rows": count, "out": args.out}))
        return 0
    except (SourceError, OSError, sqlite3.Error) as exc:
        print(json.dumps({"error": str(exc)}))
        return (
            2
            if isinstance(exc, SourceError)
            and exc.code in ("incomplete_run", "no_reconciled_run")
            else 1
        )


if __name__ == "__main__":
    raise SystemExit(main())
