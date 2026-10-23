"""Text-safe decision workbooks with retained source and review evidence."""

import json
import re
import sqlite3
from collections.abc import Iterable
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from pydantic import TypeAdapter

from core.matching.contracts import InputItem
from src.matching.engine import EngineResult
from src.matching.releases import Release
from src.updater.reader import Manifest

MAX_ROWS = 1_048_576


def write_report(path: Path, rows: Iterable[tuple[InputItem, EngineResult, str]], *, release: Release, manifest: Manifest, observed_at: datetime, mode: str, connection: sqlite3.Connection) -> None:
    book = Workbook(write_only=True)
    sheets = {}
    counts = {}
    parts = {}
    headers = {
        '판정요약': ('decision ID', '입력행', '상태', '모드', '환경', '사유', '선택후보', '요청수', '부족수', '검색완료', '평가완료'),
        '검토필요': ('decision ID', '입력행', '상태', '사유', '후보', 'evidence hash', 'human resolution', 'pending approval'),
        '판정근거': ('decision ID', '후보', '원본품명', '원본규격', '필드', '결과', '사유', '기대', '관측', '확률', '순위점수', 'route'),
        '데이터버전': ('키', '값'),
    }

    def append(group: str, values: tuple[str | float | int | bool | Decimal | None, ...], header: tuple[str, ...]) -> None:
        if group not in sheets or counts[group] >= MAX_ROWS:
            parts[group] = parts.get(group, 0) + 1
            suffix = '' if parts[group] == 1 else f'_{parts[group]}'
            base = re.sub(r'[\\/*?:\[\]]', '_', group) or '업체'
            title = base[:31-len(suffix)] + suffix
            while title.casefold() in {sheet.title.casefold() for sheet in book.worksheets}:
                parts[group] += 1
                suffix = f'_{parts[group]}'
                title = base[:31-len(suffix)] + suffix
            sheet = book.create_sheet(title)
            sheets[group] = sheet
            counts[group] = 0
            cells = []
            for value in header:
                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = 's'
                cell.font = Font(bold=True)
                cells.append(cell)
            sheet.append(cells)
            counts[group] = 1
        sheet = sheets[group]
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=float(value) if isinstance(value, Decimal) else value)
            if isinstance(value, str):
                cell.data_type = 's'
            if value == '자동선택':
                cell.fill = PatternFill('solid', fgColor='C6EFCE')
            elif value in ('review_required', 'shadow', 'fixture'):
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
            cells.append(cell)
        sheet.append(cells)
        counts[group] += 1

    for item, result, job in rows:
        decision = result.decision
        reasons = ','.join(reason.value for reason in decision.reasons)
        append('판정요약', (job, item.input_row_id, decision.status.value, mode, decision.environment, reasons, ','.join(decision.candidates), result.requested_count, result.shortage, result.retrieval.scope_scan_complete, result.retrieval.decision_evaluation_complete), headers['판정요약'])
        group = '업체_' + (item.sheet or item.seller_name or '미상')
        comparison_header = ('decision ID', '입력 원본품명', '입력 원본규격', '희망가', '입력가격종류', '입력단위', '후보', '후보 원본품명', '후보 원본규격', '현재단가', '원가정확값', '가격종류', '통화', '단위', '세금', '관측시각', '요청판정', '모드', '환경', '후보선택', '후보필수조건', '후보사유')
        checks = result.candidate_checks
        if not checks:
            append(group, (job, item.name_raw, item.spec_raw, item.price.amount, item.price.kind, item.price.unit, '', '', '', None, '', '', '', '', '', observed_at.isoformat(), decision.status.value, mode, decision.environment, '후보없음', '', ''), comparison_header)
            if decision.status.value in ('review_required', 'invalid_input', 'processing_failed'):
                append('검토필요', (job, item.input_row_id, decision.status.value, reasons, '', '', '', ''), headers['검토필요'])
        for check in checks:
            candidate = check.candidate
            from core.matching.constraints import aggregate_status
            selected = candidate.listing_key in decision.candidates
            selection = '미선택'
            if selected:
                selection = '자동선택' if mode == 'auto' and decision.status.value == 'auto_confirmed' else '검토후보' if mode == 'auto' else 'shadow 선택'
            candidate_status = aggregate_status(check.constraints).value
            candidate_reasons = ','.join(c.reason.value for c in check.constraints if c.status.value != 'PASS')
            append(group, (job, item.name_raw, item.spec_raw, item.price.amount, item.price.kind, item.price.unit, candidate.listing_key, candidate.name_raw, candidate.spec_raw, candidate.price.amount, str(candidate.price.amount), candidate.price.kind, candidate.price.currency, candidate.price.unit, candidate.price.tax_basis, json.dumps({key: value.isoformat() if value else None for key, value in candidate.group_checked_at}), decision.status.value, mode, candidate.environment, selection, candidate_status, candidate_reasons), comparison_header)
            evidence = connection.execute('SELECT evidence_hash FROM candidate_checks WHERE job_id=? AND candidate_id=?', (job, candidate.listing_key)).fetchone()
            review = connection.execute('SELECT resolution_status,pending_approval FROM review_events WHERE job_id=? AND candidate_id=? ORDER BY rowid DESC LIMIT 1', (job, candidate.listing_key)).fetchone()
            if decision.status.value in ('review_required', 'invalid_input', 'processing_failed'):
                append('검토필요', (job, item.input_row_id, decision.status.value, reasons, candidate.listing_key, evidence[0], review[0] if review else '', bool(review and review[1])), headers['검토필요'])
            for constraint in check.constraints:
                append('판정근거', (job, candidate.listing_key, candidate.name_raw, candidate.spec_raw, constraint.field, constraint.status.value, constraint.reason.value, str(constraint.expected), str(constraint.actual), check.suitability_probability, check.rank_score, check.route), headers['판정근거'])
    for name, header in headers.items():
        if name not in sheets:
            append(name, (), header)
    manifest_data = json.loads(TypeAdapter(Manifest).dump_json(manifest))
    for key, value in {'mode': mode, 'report_time': observed_at.isoformat(), 'release_hash': release.content_hash, 'release': release.model_dump(mode='json'), 'source': manifest_data, 'business_accuracy': 'not measured by this report'}.items():
        append('데이터버전', (key, json.dumps(value, ensure_ascii=False)), headers['데이터버전'])
    # Exclusive file creation prevents silently replacing a prior decision workbook.
    with path.open('xb') as stream:
        book.save(stream)
