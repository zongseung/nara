import tempfile
import unittest
from pathlib import Path
from decimal import Decimal
from openpyxl import Workbook
from core.matching.contracts import Purpose, EvidenceState, ContractError
from core.matching.normalization import normalize_dimensions, normalize_material
from src.format_convert import read_input_items


class InputTests(unittest.TestCase):
    def test_equivalent_units(self):
        self.assertEqual(
            normalize_dimensions("20x30cm").value,
            normalize_dimensions("200×300mm").value,
        )

    def test_three_axes(self):
        self.assertEqual(
            normalize_dimensions("20x30x0.12cm").value,
            (Decimal("200"), Decimal("300"), Decimal("1.20")),
        )

    def test_no_unit_not_known(self):
        self.assertEqual(
            normalize_dimensions("200x300").status, EvidenceState.AMBIGUOUS
        )

    def test_dictionaryless_dollar(self):
        self.assertEqual(
            normalize_dimensions("200$300$1.2").status, EvidenceState.AMBIGUOUS
        )

    def test_material_grade_separate(self):
        self.assertNotEqual(
            normalize_material("SUS304").value, normalize_material("SUS316").value
        )

    def test_composite_unknown(self):
        self.assertEqual(normalize_material("ABS/PVC").status, EvidenceState.AMBIGUOUS)

    def test_xlsx_origin_and_missing_price(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "input.xlsx"
            w = Workbook()
            s = w.active
            s.title = "not-an-affiliate"
            s.append(["물품식별번호", "한글품명", "재질"])
            s.append(["00123", "판 20x30cm", "SUS304"])
            w.save(p)
            w.close()
            item = read_input_items(
                p, purpose=Purpose.EXACT_IDENTITY, category_id="fixture"
            )[0]
            self.assertEqual(
                (
                    item.goods_id,
                    item.row,
                    item.sheet,
                    item.price.amount,
                    item.seller_id,
                ),
                ("00123", 2, "not-an-affiliate", None, None),
            )
            self.assertEqual(len(item.file_sha256), 64)

    def test_formula_price_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "bad.xlsx"
            w = Workbook()
            s = w.active
            s.append(["물품식별번호", "한글품명", "희망가(부가세포함)"])
            s.append(["001", "item", "=1+1"])
            w.save(p)
            w.close()
            with self.assertRaises(ContractError):
                read_input_items(
                    p, purpose=Purpose.EXACT_IDENTITY, category_id="fixture"
                )

    def test_actual_xls_and_xlsx_same_values(self):
        import struct
        import xlrd

        def record(code, data):
            return struct.pack("<HH", code, len(data)) + data

        rows = [
            ["물품식별번호", "한글품명", "재질", "희망가(부가세포함)"],
            ["00123", "판 20x30cm", "SUS304", "1,300,000"],
        ]
        payload = record(9, struct.pack("<HH", 2, 16)) + record(
            0x42, struct.pack("<H", 949)
        )
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                raw = value.encode("cp949")
                payload += record(
                    4, struct.pack("<HH", i, j) + b"\0\0\0" + bytes([len(raw)]) + raw
                )
        payload += record(10, b"")
        with tempfile.TemporaryDirectory() as d:
            old = Path(d) / "input.xls"
            old.write_bytes(payload)
            new = Path(d) / "input.xlsx"
            w = Workbook()
            s = w.active
            for row in rows:
                s.append(row)
            w.save(new)
            w.close()
            a = read_input_items(
                old, purpose=Purpose.EXACT_IDENTITY, category_id="fixture"
            )[0]
            b = read_input_items(
                new, purpose=Purpose.EXACT_IDENTITY, category_id="fixture"
            )[0]
            self.assertEqual(
                (a.goods_id, a.name_raw, a.attributes, a.price),
                (b.goods_id, b.name_raw, b.attributes, b.price),
            )

    def test_duplicate_header_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "bad.xlsx"
            w = Workbook()
            s = w.active
            s.append(["물품식별번호", "한글품명", "한글품명"])
            s.append(["001", "a", "b"])
            w.save(p)
            w.close()
            with self.assertRaises(ContractError):
                read_input_items(
                    p, purpose=Purpose.EXACT_IDENTITY, category_id="fixture"
                )

    def test_conflicting_units(self):
        self.assertEqual(
            normalize_dimensions("20cm x 300mm").status, EvidenceState.CONFLICTING
        )

    def test_unsupported_unit_not_silently_mm(self):
        self.assertEqual(
            normalize_dimensions("20x30mmHg").status, EvidenceState.AMBIGUOUS
        )
