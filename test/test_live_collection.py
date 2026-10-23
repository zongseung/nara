import json
import io
import os
from contextlib import redirect_stdout
from dataclasses import replace
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import tempfile
from threading import Event, Lock, Thread
import unittest
from multiprocessing import get_context
from unittest.mock import patch
from openpyxl import load_workbook

from core.catalog_sources import SearchQuery, SourceError, get_profile
from src.live_collection import collect_live, export_live, ProcessRequestLimiter
from test.test_transport_limits import Clock
from test.updater_fixtures import row


class LiveCollectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.products = [dict(row(str(i), '100.000000000000000001'), totCnt=250,
                              extra_field='keep me') for i in range(250)]
        self.calls = []
        self.bad_page = None
        self.status = 200
        self.active = set()
        self.parallel = Event()
        self.guard = Lock()
        owner = self

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                query = json.loads(self.rfile.read(int(self.headers['Content-Length'])))['searchVO']
                number = int(query['currentPage'])
                with owner.guard:
                    owner.calls.append(query)
                    owner.active.add(number)
                    if 2 in owner.active and 3 in owner.active:
                        owner.parallel.set()
                if number in (2,3) and owner.bad_page is None:
                    owner.parallel.wait(8)
                start = (number-1)*100
                products = owner.products[start:start+100]
                if number == owner.bad_page:
                    products = []
                payload = json.dumps(dict(totalSize=400, rsltList=products)).encode()
                self.send_response(owner.status)
                self.send_header('Content-Length', str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)
                with owner.guard:
                    owner.active.discard(number)

            def log_message(self, *args):
                return

        self.server = ThreadingHTTPServer(('127.0.0.1',0), Handler)
        self.thread = Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.addCleanup(self.stop_server)
        self.profile = replace(get_profile('shopping-all'), endpoint=f'http://127.0.0.1:{self.server.server_port}/search')

    def stop_server(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()

    def test_pages_use_separate_processes_and_preserve_order(self):
        result = collect_live(SearchQuery(keyword='sample'), limit=230, workers=2, profile=self.profile)
        self.assertEqual(len(result.items), 230)
        self.assertEqual([item.listing_key.rsplit(':',1)[-1] for item in result.items], [str(i) for i in range(230)])
        self.assertTrue(self.parallel.is_set())
        self.assertGreaterEqual(len(set(result.process_ids)), 2)
        self.assertNotIn(os.getpid(), result.process_ids)
        self.assertEqual(result.matched_total, 250)
        self.assertEqual(result.response_total, 400)
        self.assertEqual(len(self.calls), 3)
        self.assertTrue(all(query['searchKeyword']=='sample' for query in self.calls))
        self.assertFalse((self.root/'control').exists())

    def test_short_page_does_not_publish_a_success_file(self):
        self.bad_page = 2
        destination = self.root/'failed.jsonl'
        with self.assertRaises(SourceError):
            export_live(SearchQuery(), destination, limit=200, workers=2, profile=self.profile)
        self.assertFalse(destination.exists())
        self.assertEqual(list(self.root.glob('*.partial')), [])

    def test_duplicate_item_does_not_publish_a_success_file(self):
        self.products[100] = self.products[0]
        destination = self.root/'duplicate.jsonl'
        with self.assertRaises(SourceError) as failure:
            export_live(SearchQuery(), destination, limit=230, workers=2, profile=self.profile)
        self.assertEqual(failure.exception.code, 'live_duplicate_item')
        self.assertFalse(destination.exists())

    def test_jsonl_preserves_raw_fields_and_marks_truncation(self):
        destination = self.root/'result.jsonl'
        report = export_live(SearchQuery(keyword='sample'), destination, limit=10, workers=2, profile=self.profile)
        lines = [json.loads(line) for line in destination.read_text().splitlines()]
        self.assertTrue(report['truncated'])
        self.assertFalse(report['source_snapshot_verified'])
        self.assertEqual(report['collected_items'], 10)
        self.assertEqual(lines[1]['raw']['extra_field'], 'keep me')
        self.assertEqual(lines[1]['raw']['ctrtUprc'], '100.000000000000000001')
        self.assertEqual(len(lines), 11)
        self.assertEqual(len(self.calls), 1)
        with self.assertRaises(SourceError):
            export_live(SearchQuery(), destination, profile=self.profile)
        self.assertEqual(len(self.calls), 1)

    def test_invalid_settings_do_not_access_source(self):
        for limit,workers in ((0,2),(10001,2),(1,0),(1,5)):
            with self.assertRaises(SourceError):
                collect_live(SearchQuery(), limit=limit, workers=workers, profile=self.profile)
        for minimum,maximum in (('1',''),('','2'),('a','2'),('-1','2'),('3','2'),('NaN','2'),('0','Infinity')):
            with self.assertRaises(SourceError):
                collect_live(SearchQuery(start_price=minimum,end_price=maximum), profile=self.profile)
        self.assertEqual(self.calls, [])

    def test_cli_workbook_has_literal_values_without_creating_a_database(self):
        from src.updater.cli import main
        destination = self.root/'live.xlsx'
        data = self.root/'unused-store'
        self.products[0].update(itemIdnfNm='=2+2',itemIndfNmView='=2+2')
        output = io.StringIO()
        with patch('sys.argv', ['nara','--data-dir',str(data),'collect','--keyword','=1+1',
                               '--company','seller','--min-price','100','--max-price','100',
                               '--limit','10','--workers','2','--out',str(destination)]), \
                patch('src.live_collection.get_profile', return_value=self.profile), redirect_stdout(output):
            main()
        self.assertEqual(json.loads(output.getvalue())['collected_items'], 10)
        self.assertFalse(data.exists())
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(self.calls[0]['etpsNm'], 'seller')
        self.assertEqual(self.calls[0]['startPrce'], '100')
        self.assertEqual(self.calls[0]['endPrce'], '100')
        book = load_workbook(destination, read_only=True, data_only=False)
        self.addCleanup(book.close)
        metadata = {name.value:value for name,value in book['수집정보'].iter_rows(min_row=2)}
        self.assertEqual(metadata['keyword'].value, '=1+1')
        self.assertEqual(metadata['keyword'].data_type, 's')
        self.assertEqual(book['실시간품목']['C2'].value, '=2+2')
        self.assertEqual(book['실시간품목']['C2'].data_type, 's')
        self.assertEqual(len(list(book['실시간품목'].values)), 11)

    def test_http_failure_keeps_its_status_across_process_boundary(self):
        self.status = 401
        with self.assertRaises(SourceError) as failure:
            collect_live(SearchQuery(), limit=100, workers=2, profile=self.profile)
        self.assertEqual(failure.exception.code, 'http_status')
        self.assertEqual(failure.exception.status, 401)


class ProcessBudgetTests(unittest.TestCase):
    def test_request_spacing_and_cooldown_are_shared(self):
        clock = Clock()
        with get_context('spawn').Manager() as manager, \
                patch('src.live_collection.monotonic', clock.read), \
                patch('src.live_collection.sleep', clock.sleep):
            gate = manager.Lock()
            next_start = manager.Value('d',0.0)
            first = ProcessRequestLimiter(gate,next_start)
            second = ProcessRequestLimiter(gate,next_start)
            first.before_request()
            second.before_request()
            self.assertEqual(clock.now, 1.0)
            first.pause_all(5)
            second.before_request()
            self.assertEqual(clock.now, 6.0)
            first.before_request()
            self.assertEqual(clock.now, 7.0)
