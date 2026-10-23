import io
import json
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from dataclasses import replace
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from threading import Thread
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

from core.catalog_sources import SourceError, get_profile
from src.build_report import build
from src.compare_file import main
from src.transport import RequestLimiter
from test.updater_fixtures import row


class CompanyComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.out = Path(self.temporary.name)/'comparison.xlsx'
        self.calls = []
        self.image_calls = 0
        self.only_other_company = False
        self.candidates = None
        self.product_overrides = {}
        png = io.BytesIO()
        Image.new('RGB', (80,60), 'green').save(png,format='PNG')
        picture = png.getvalue()
        self.picture = picture
        owner = self

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                query = json.loads(self.rfile.read(int(self.headers['Content-Length'])))['searchVO']
                owner.calls.append(query)
                name = query['etpsNm']
                if name == '지피코리아' and not owner.only_other_company:
                    products = [owner.product('ours1','주식회사 지피코리아'),
                                owner.product('ours2','주식회사 지피코리아'),
                                owner.product('other','지피코리아 기타')]
                elif name == '회사B':
                    products = [owner.product('b1','주식회사 회사B')]
                elif name:
                    products = [owner.product('other','다른업체')]
                else:
                    products = [owner.product('self1','주식회사 지피코리아'),
                                owner.product('self2','주식회사 회사B')]
                    products.extend(dict(owner.product(f'competitor{i}',f'경쟁사{i}'),ctrtUprc=str(101+i))
                                    for i in range(3))
                    if owner.candidates is not None:
                        products = owner.candidates
                for product in products:
                    product['totCnt'] = len(products)
                payload = json.dumps({'totalSize':len(products),'rsltList':products}).encode()
                self.send_response(200)
                self.send_header('Content-Length',str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)

            def do_GET(self):
                owner.image_calls += 1
                self.send_response(200)
                self.send_header('Content-Type','image/png')
                self.send_header('Content-Length',str(len(picture)))
                self.end_headers()
                self.wfile.write(picture)

            def log_message(self, *args):
                return

        self.server = ThreadingHTTPServer(('127.0.0.1',0),Handler)
        thread = Thread(target=self.server.serve_forever,daemon=True)
        thread.start()
        self.addCleanup(thread.join)
        self.addCleanup(self.server.server_close)
        self.addCleanup(self.server.shutdown)
        self.base = f'http://127.0.0.1:{self.server.server_port}'
        profile = replace(get_profile('shopping-all'),endpoint=self.base+'/search')
        self.enterContext(patch('src.live_collection.get_profile',return_value=profile))
        self.enterContext(patch('src.g2b_client.get_profile',return_value=profile))
        self.enterContext(patch('src.g2b_client._DEFAULT_LIMITER',RequestLimiter(rate=10000)))

    def product(self, identity, company):
        product = dict(row(identity,'100'),mnftrEtpsNm=company,ctentUntyGrpNm=company,
                    itemIdnfNo=identity,itemIdnfNm=f'안내판, GK-{identity}, 600×150mm',
                    dtlsPrnmNo='5512171801',pdctAtrbCdDtlNm='재질:목재,용도:종합안내,지지대수:0개',
                    imgSrc=self.base+'/image.png')
        product.update(self.product_overrides.get(identity, {}))
        return product

    def test_company_names_produce_real_company_sheets_and_embedded_photos_without_input(self):
        with patch('sys.argv',['compare_file.py','--company','지피코리아','--company','회사B',
                               '--keyword','안내판','--limit','3','--workers','2','--out',str(self.out)]), \
                redirect_stdout(io.StringIO()):
            main()
        book = load_workbook(self.out)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames,['주식회사 지피코리아','주식회사 회사B'])
        sheet = book.worksheets[0]
        self.assertEqual(sheet['D4'].value,'ours1')
        self.assertEqual(sheet['D12'].value,'ours2')
        self.assertEqual(sheet['D3'].fill.fgColor.rgb,'00FFFF00')
        self.assertEqual({sheet.cell(3,column).value for column in (5,6,7)}, {'경쟁사0','경쟁사1','경쟁사2'})
        self.assertEqual([sheet.cell(9,column).value for column in (4,5,6,7)], ['100','101','102','103'])
        with ZipFile(self.out) as archive:
            self.assertEqual(len([name for name in archive.namelist() if name.startswith('xl/media/')]),12)
        self.assertEqual(self.image_calls,12)
        self.assertTrue(all(query['searchKeyword']=='안내판' for query in self.calls if query['etpsNm']))
        self.assertTrue(all(query['searchKeyword']=='안내판 종합안내' for query in self.calls if not query['etpsNm']))
        self.assertFalse((self.out.parent/'storage').exists())

    def test_company_filter_never_uses_a_different_company_as_our_items(self):
        self.only_other_company = True
        with patch('sys.argv',['compare_file.py','--company','지피코리아','--out',str(self.out)]), \
                redirect_stdout(io.StringIO()), self.assertRaises(SourceError) as failure:
            main()
        self.assertEqual(failure.exception.code,'company_not_found')
        self.assertFalse(self.out.exists())
        self.assertEqual(len(self.calls),1)

    def test_help_displays_percentage_without_any_source_access(self):
        stdout = io.StringIO()
        with patch('sys.argv',['compare_file.py','--help']), redirect_stdout(stdout), \
                self.assertRaises(SystemExit) as exit_status:
            main()
        self.assertEqual(exit_status.exception.code, 0)
        self.assertIn('10%', stdout.getvalue())
        self.assertEqual(self.calls, [])

    def test_invalid_sources_and_settings_fail_before_any_source_access(self):
        for options in ([],['input.xls','--company','지피코리아'],['--company',' '],
                        ['--company','지피코리아','--company','주식회사 지피코리아'],
                        ['--company','지피코리아','--mode','auto'],
                        ['--company','지피코리아','--catalog-db','missing.db'],
                        ['--company','지피코리아','--limit','-1'],
                        ['--company','지피코리아','--limit','10001'],
                        ['--company','지피코리아','--top','0']):
            with self.subTest(options=options), patch('sys.argv',['compare_file.py',*options,'--out',str(self.out)]), \
                    redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as failure:
                main()
            self.assertEqual(failure.exception.code,2)
        self.assertEqual(self.calls,[])

    def test_direction_material_and_geometry_select_real_candidates_then_sort_columns_by_price(self):
        self.product_overrides['ours1'] = {'itemIdnfNm':'안내판, OUR1, 600×150mm, 양면',
                                         'pdctAtrbCdDtlNm':'재질:목재,용도:1방향안내,지지대수:0개'}
        self.candidates = []
        for identity, price, size, material, use, face, category in (
                ('expensive',130,'560×140mm','목재','1방향안내','양면','5512171801'),
                ('cheap',110,'550×140mm','목재','1방향안내','양면','5512171801'),
                ('mixed',120,'570×145mm','목재/스테인리스','1방향안내','양면','5512171801'),
                ('wrong_direction',105,'550×140mm','목재','2방향안내','양면','5512171801'),
                ('wrong_face',105,'550×140mm','목재','1방향안내','단면','5512171801'),
                ('accessory',105,'550×140mm','목재','1방향안내','양면','5512190801'),
                ('same_price',100,'550×140mm','목재','1방향안내','양면','5512171801'),
                ('too_cheap',90,'550×140mm','목재','1방향안내','양면','5512171801')):
            item = self.product(identity, identity)
            item.update(ctrtUprc=str(price),itemIdnfNm=f'안내판, {identity}, {size}, {face}',
                        pdctAtrbCdDtlNm=f'재질:{material},용도:{use},지지대수:0개',dtlsPrnmNo=category)
            self.candidates.append(item)
        with patch('sys.argv',['compare_file.py','--company','지피코리아','--limit','1',
                               '--workers','1','--out',str(self.out)]), redirect_stdout(io.StringIO()):
            main()
        book = load_workbook(self.out)
        self.addCleanup(book.close)
        sheet = book.worksheets[0]
        self.assertEqual([sheet.cell(9,c).value for c in (5,6,7)], ['110','120','130'])
        self.assertEqual([sheet.cell(3,c).value for c in (5,6,7)], ['cheap','mixed','expensive'])
        self.assertEqual(sheet['F7'].value, '목재/스테인리스')
        self.assertIn('양면', sheet['D8'].value)
        self.assertIn('혼합', sheet['F3'].comment.text)
        with ZipFile(self.out) as archive:
            pictures = [archive.read(name) for name in archive.namelist() if name.startswith('xl/media/')]
            self.assertEqual(pictures, [self.picture]*4)

    def test_existing_output_is_preserved_before_any_source_access(self):
        self.out.write_bytes(b'original')
        for source in (['--company','지피코리아'], ['input.xls']):
            with self.subTest(source=source), \
                    patch('sys.argv',['compare_file.py',*source,'--out',str(self.out)]), \
                    redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as failure:
                main()
            self.assertEqual(failure.exception.code, 2)
        self.assertEqual(self.out.read_bytes(), b'original')
        self.assertEqual(self.calls, [])

    def test_missing_output_directories_are_created_and_workbook_is_readable(self):
        self.out = self.out.parent/'nested'/'한글 폴더'/'comparison.XLSX'
        with patch('sys.argv',['compare_file.py','--company','지피코리아','--limit','1',
                               '--top','1','--workers','1','--out',str(self.out)]), \
                redirect_stdout(io.StringIO()):
            main()
        with self.out.open('rb') as stream:
            book = load_workbook(stream)
            self.addCleanup(book.close)
            self.assertEqual(book.worksheets[0]['D4'].value, 'ours1')
            self.assertEqual(book.worksheets[0]['E4'].value, 'competitor0')
        with ZipFile(self.out) as archive:
            self.assertEqual(len([name for name in archive.namelist() if name.startswith('xl/media/')]), 2)
        self.assertEqual(list(self.out.parent.iterdir()), [self.out])

    def test_invalid_extension_identifies_output_before_any_source_access(self):
        output = self.out.parent/'missing'/'comparison.xslx'
        for source in (['--company','지피코리아'], ['input.xls']):
            stderr = io.StringIO()
            with self.subTest(source=source), \
                    patch('sys.argv',['compare_file.py',*source,'--out',str(output)]), \
                    redirect_stderr(stderr), self.assertRaises(SystemExit) as failure:
                main()
            self.assertEqual(failure.exception.code, 2)
            self.assertIn('.xlsx', stderr.getvalue())
            self.assertIn(str(output), stderr.getvalue())
        self.assertEqual(self.calls, [])
        self.assertFalse(output.parent.exists())

    def test_file_blocking_output_directory_fails_before_any_source_access(self):
        parent = self.out.parent/'not-a-directory'
        parent.write_bytes(b'original')
        for source in (['--company','지피코리아'], ['input.xls']):
            stderr = io.StringIO()
            with self.subTest(source=source), \
                    patch('sys.argv',['compare_file.py',*source,'--out',str(parent/'comparison.xlsx')]), \
                    redirect_stderr(stderr), self.assertRaises(SystemExit) as failure:
                main()
            self.assertEqual(failure.exception.code, 2)
            self.assertIn(str(parent), stderr.getvalue())
        self.assertEqual(parent.read_bytes(), b'original')
        self.assertEqual(self.calls, [])

    def test_python_builder_creates_output_directory_and_preserves_existing_file(self):
        output = self.out.parent/'nested'/'comparison.xlsx'
        build([], output)
        book = load_workbook(output)
        book.close()
        original = output.read_bytes()
        with self.assertRaises(FileExistsError):
            build([], output)
        self.assertEqual(output.read_bytes(), original)

    def test_comparison_requires_matching_supports_and_each_dimension_within_ten_percent(self):
        self.product_overrides['ours1'] = {'pdctAtrbCdDtlNm':'재질:목재,용도:종합안내,지지대수:1개'}
        self.candidates = []
        for identity, size, supports, price in (
                ('minimum', '540×135mm', '1개', 110),
                ('maximum', '660×165mm', '1개', 130),
                ('nearby', '570×145mm', '1개', 120),
                ('too_wide', '660.01×150mm', '1개', 101),
                ('too_short', '600×134.99mm', '1개', 102),
                ('too_small', '100×100mm', '1개', 103),
                ('two_legs', '600×150mm', '2개', 104),
                ('no_legs', '600×150mm', '0개', 105),
                ('unknown_legs', '600×150mm', None, 106)):
            item = self.product(identity, identity)
            item.update(itemIdnfNm=f'안내판, {identity}, {size}', ctrtUprc=str(price),
                        pdctAtrbCdDtlNm='재질:목재,용도:종합안내'+
                            (f',지지대수:{supports}' if supports is not None else ''))
            self.candidates.append(item)
        with patch('sys.argv',['compare_file.py','--company','지피코리아','--limit','1',
                               '--workers','1','--out',str(self.out)]), redirect_stdout(io.StringIO()):
            main()
        book = load_workbook(self.out)
        self.addCleanup(book.close)
        sheet = book.worksheets[0]
        self.assertEqual([sheet.cell(3,c).value for c in (5,6,7)], ['minimum','nearby','maximum'])
        self.assertIn('지지대 1개', sheet['D8'].value)

    def test_unknown_reference_support_count_defers_comparison_with_visible_reason(self):
        self.product_overrides['ours1'] = {'pdctAtrbCdDtlNm':'재질:목재,용도:종합안내'}
        with patch('sys.argv',['compare_file.py','--company','지피코리아','--limit','1',
                               '--workers','1','--out',str(self.out)]), redirect_stdout(io.StringIO()):
            main()
        book = load_workbook(self.out)
        self.addCleanup(book.close)
        sheet = book.worksheets[0]
        self.assertIsNone(sheet['E4'].value)
        self.assertIn('지지대', sheet['E3'].value)
        self.assertIn('미확인', sheet['E3'].value)
        self.assertFalse(any(not query['etpsNm'] for query in self.calls))

    def test_exact_price_does_not_bypass_support_and_dimension_checks(self):
        self.product_overrides['ours1'] = {'pdctAtrbCdDtlNm':'재질:목재,용도:종합안내,지지대수:0개'}
        self.candidates = []
        for identity, size, supports in (('two_legs','600×150mm',2),
                                         ('oversized','800×255mm',0),
                                         ('valid','600×150mm',0)):
            item = self.product(identity, identity)
            item.update(itemIdnfNm=f'안내판, {identity}, {size}',
                        pdctAtrbCdDtlNm=f'재질:목재,용도:종합안내,지지대수:{supports}개')
            self.candidates.append(item)
        with patch('sys.argv',['compare_file.py','--company','지피코리아','--limit','1',
                               '--exact-price','--fallback-band','0','--workers','1',
                               '--out',str(self.out)]), redirect_stdout(io.StringIO()):
            main()
        book = load_workbook(self.out)
        self.addCleanup(book.close)
        sheet = book.worksheets[0]
        self.assertEqual(sheet['E3'].value, 'valid')
        self.assertIsNone(sheet['F4'].value)
        self.assertIsNone(sheet['G4'].value)

    def test_approximate_price_retry_keeps_structure_and_local_price_bounds(self):
        self.candidates = []
        for identity, size, supports, price in (('two_legs','600×150mm',2,105),
                                                ('oversized','800×255mm',0,105),
                                                ('outside_price_band','600×150mm',0,150),
                                                ('valid','600×150mm',0,110)):
            item = self.product(identity, identity)
            item.update(itemIdnfNm=f'안내판, {identity}, {size}', ctrtUprc=str(price),
                        pdctAtrbCdDtlNm=f'재질:목재,용도:종합안내,지지대수:{supports}개')
            self.candidates.append(item)
        with patch('sys.argv',['compare_file.py','--company','지피코리아','--limit','1',
                               '--exact-price','--workers','1','--out',str(self.out)]), \
                redirect_stdout(io.StringIO()):
            main()
        book = load_workbook(self.out)
        self.addCleanup(book.close)
        sheet = book.worksheets[0]
        self.assertEqual(sheet['E4'].value, 'valid')
        self.assertIn('근사', sheet['E3'].value)
        self.assertIsNone(sheet['F4'].value)

    def test_complete_candidates_survive_detail_budget_and_smaller_higher_price_wins(self):
        self.product_overrides['ours1'] = {'itemIdnfNm':'안내판, OUR1, 600×150mm, 양면',
                                         'pdctAtrbCdDtlNm':'재질:목재,용도:1방향안내,지지대수:0개'}
        self.candidates = []
        for index in range(24):
            item = self.product(f'missing{index}',f'미기재{index}')
            item['pdctAtrbCdDtlNm'] = '재질:목재,용도:1방향안내,지지대수:0개'
            self.candidates.append(item)
        for identity, price, size in (('larger',110,'650×160mm'),('smaller',130,'570×140mm')):
            item = self.product(identity, identity)
            item.update(ctrtUprc=str(price),itemIdnfNm=f'안내판, {identity}, {size}, 양면',
                        pdctAtrbCdDtlNm='재질:목재,용도:1방향안내,지지대수:0개')
            self.candidates.append(item)
        with patch('sys.argv',['compare_file.py','--company','지피코리아','--limit','1','--top','1',
                               '--workers','1','--out',str(self.out)]), redirect_stdout(io.StringIO()):
            main()
        book = load_workbook(self.out)
        self.addCleanup(book.close)
        sheet = book.worksheets[0]
        self.assertEqual((sheet['E3'].value,sheet['E9'].value),('smaller','130'))
