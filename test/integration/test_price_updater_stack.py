"""Explicit isolated Docker/Prefect fixture QA; run only after QA Compose is started."""
import json
import os
from uuid import uuid4
from pathlib import Path
import subprocess
import unittest
from urllib.error import HTTPError
from urllib.request import urlopen, Request

EVIDENCE=Path('.omo/evidence/prefect-operations')
COMPOSE=['docker','compose','--env-file',os.environ.get('NARA_QA_ENV_FILE','/tmp/nara-updater-qa.env'),'-p','nara-updater-qa','-f','docker/compose.yaml','-f','docker/compose.fixture.yaml']


def command(args):
    result=subprocess.run(COMPOSE+args,capture_output=True,text=True,timeout=180)
    if result.returncode:
        raise RuntimeError(result.stderr+result.stdout)
    return result.stdout


@unittest.skipUnless(os.environ.get('NARA_STACK_QA')=='1','explicit isolated Compose QA only')
class StackTests(unittest.TestCase):
    def test_fixture_stack(self):
        with self.assertRaises(HTTPError) as error:
            urlopen(Request('http://127.0.0.1:14200/api/deployments/filter',data=b'{}',headers={'Content-Type':'application/json'}),timeout=5)
        self.assertEqual(error.exception.code,401)
        registrations=[]
        for _ in range(2):
            registrations.append(command(['exec','-T','updater-worker','python','-m','src.updater.bootstrap']).strip().splitlines()[-1])
        self.assertEqual(registrations[0],registrations[1])
        runs=[]
        data_root="/data/qa-"+uuid4().hex
        (EVIDENCE/'stack-data-root.txt').write_text(data_root)
        script='''
import json,time
from pathlib import Path
from datetime import datetime,timezone
from prefect.client.orchestration import get_client
from prefect.client.schemas.filters import FlowRunFilter,FlowRunFilterId
from src.updater.dispatcher import ensure_slot,submit_pending
from src.updater.store import Store
store=Store(Path(DATA_ROOT))
key=ensure_slot(store,datetime.now(timezone.utc),manual=True)
run_id=submit_pending(store,key,parameters=dict(environment='fixture',source_endpoint='http://fixture-api:8080/search',data_root=DATA_ROOT))
if run_id is None: raise RuntimeError('submission held')
from uuid import UUID
with get_client(sync_client=True) as client:
    for count in range(70):
        run=client.read_flow_run(UUID(run_id))
        if run.state.is_final():
            tasks=client.read_task_runs(flow_run_filter=FlowRunFilter(id=FlowRunFilterId(any_=[run.id])))
            print(json.dumps(dict(run_id=run_id,state=run.state.type.value,slot_key=key,tasks=[dict(name=t.name,run_count=t.run_count) for t in tasks])))
            break
        time.sleep(2)
    else: raise RuntimeError('fixture run timeout')
'''
        script=script.replace('DATA_ROOT',repr(data_root))
        for price in ('100','100.0','120','100'):
            command(['exec','-T','fixture-api','python','-c',"from pathlib import Path; Path('/fixture/state.json').write_text("+repr(json.dumps(dict(price=price)))+")"])
            output=command(['exec','-T','updater-worker','python','-c',script])
            run=json.loads(output.strip().splitlines()[-1])
            runs.append(run)
            self.assertEqual(run['state'],'COMPLETED',output)
        command(['exec','-T','fixture-api','python','-c',"from pathlib import Path; Path('/fixture/state.json').write_text('{\"failure\":true,\"price\":\"500\"}')"])
        failed=json.loads(command(['exec','-T','updater-worker','python','-c',script]).strip().splitlines()[-1])
        self.assertEqual(failed['state'],'FAILED')
        runs.append(failed)
        current=command(['exec','-T','updater-worker','python','-m','src.updater.cli','--data-dir',data_root,'current'])
        self.assertEqual(json.loads(current)['contract_price_decimal'],'100')
        history=command(['exec','-T','updater-worker','python','-m','src.updater.cli','--data-dir',data_root,'history','shopping:listing:ctrtItemMngNo:fixture-A','--price-only'])
        self.assertEqual([json.loads(line)['new_amount'] for line in history.splitlines()],['100','120','100'])
        command(['exec','-T','updater-worker','python','-m','src.updater.cli','--data-dir',data_root,'export','--out','/data/qa.xlsx'])
        command(['cp','updater-worker:/data/qa.xlsx',str(EVIDENCE/'stack-output.xlsx')])
        (EVIDENCE/'stack-runs.json').write_text(json.dumps(runs,indent=2))
        (EVIDENCE/'stack-current.json').write_text(current)
        (EVIDENCE/'stack-history.jsonl').write_text(history)
        from openpyxl import load_workbook
        book=load_workbook(EVIDENCE/'stack-output.xlsx',read_only=True,data_only=False)
        self.assertEqual(book['상품1']['C2'].data_type,'s')
        self.assertEqual(book['상품1']['C2'].value,'=1+1')
        self.assertIn('가격이력1',book.sheetnames)
        book.close()
        command(['exec','-T','fixture-api','python','-c',"from pathlib import Path; Path('/fixture/state.json').write_text('{\"price\":\"100\",\"count\":101,\"page2_failures\":3}')"])
        retry=json.loads(command(['exec','-T','updater-worker','python','-c',script.replace(data_root,data_root+'-retry')]).strip().splitlines()[-1])
        self.assertEqual(retry['state'],'COMPLETED')
        self.assertTrue(any(t['name'].startswith('collect-partition') and t['run_count']==2 for t in retry['tasks']))
        (EVIDENCE/'stack-retry.json').write_text(json.dumps(retry,indent=2))
