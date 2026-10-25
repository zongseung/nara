import io
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import requests
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core.comparison import CandidatePage
from core.legacy_matching import extract
from src.web.api import create_app
from src.web.store import JobStore
from src.web.worker import execute
from test.updater_fixtures import row


def product(identifier: str, company: str, price: int, supports: str = '2개'):
    raw = row(identifier, price=str(price))
    raw.update(mnftrEtpsNm=company, itemIdnfNo=identifier, itemIdnfNm='안내판, MODEL, 600×150mm',
               pdctAtrbCdDtlNm='재질:목재,용도:방향안내,지지대수:'+supports,
               imgSrc='')
    return extract(raw)


class Provider:
    def search(self, keyword, lower='', upper='', limit=500):
        return CandidatePage(tuple(product(str(20000000+i), f'경쟁사{i}', price)['_raw']
                                   for i, price in enumerate((130, 110, 120))), 3, False, None)

    def enrich(self, raw, *, required=()):
        return raw


class WebJobsTests(unittest.TestCase):
    def __init__(self, methodName='runTest'):
        super().__init__(methodName)
        self.root = Path()
        self.sent: list[tuple[str, str]] = []
        self.store: JobStore
        self.client: TestClient

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.store = JobStore(self.root)
        self.sent = []
        self.client = TestClient(create_app(self.root, submit=lambda j, o: self.sent.append((j, o))))

    def create(self):
        response = self.client.post('/jobs', json={'company': '기준회사', 'limit': 3})
        self.assertEqual(response.status_code, 202, response.text)
        return response.json()['id']

    def collect(self, job_id, supports='2개'):
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=Provider(),
                collect=lambda request: {'기준회사': [product('10000000', '기준회사', 100, supports)]})

    def test_job_is_queued_without_collecting_inside_request(self):
        job_id = self.create()
        self.assertEqual(self.store.read(job_id).status, 'queued')
        self.assertEqual(len(self.sent), 1)
        self.assertFalse((self.root/job_id/'catalog.sqlite').exists())

    def test_research_replaces_only_requested_candidate_and_remembers_rejections(self):
        job_id = self.create()
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=Provider(),
                collect=lambda request: {'기준회사': [product('10000000', '기준회사', 100),
                                                     product('10000001', '기준회사', 100)]})
        before = self.store.item(job_id, 0)
        other = self.store.item(job_id, 1)
        self.client.post(f'/jobs/{job_id}/items/0/selection', json={'revision': 0, 'excluded': [0]})
        target = before.candidates[1]['식별번호']
        response = self.client.post(f'/jobs/{job_id}/items/0/research', json={
            'revision': 1, 'replace_item_id': target, 'conditions': {'material': '나무'}})
        self.assertEqual(response.status_code, 202, response.text)
        self.assertEqual(self.store.item(job_id, 0).candidates, before.candidates)
        self.assertEqual(response.json()['completed'], 1)
        searches = []
        class ReplacementProvider(Provider):
            def search(self, *args):
                searches.append(args)
                old = Provider().search(*args)
                return CandidatePage((*old.rows, product('30000000', '새 업체', 125)['_raw'],
                    product('30000001', before.candidates[2]['업체'], 115)['_raw']), 5, False, None)
        def unexpected(request):
            raise AssertionError('research must not recollect the company')
        execute(self.root, job_id, response.json()['operation_id'], provider=ReplacementProvider(), collect=unexpected)
        updated = self.store.item(job_id, 0)
        self.assertEqual(len(searches), 1)
        self.assertEqual(self.store.item(job_id, 1), other)
        self.assertEqual(updated.reference, before.reference)
        self.assertEqual(updated.candidates[0], before.candidates[0])
        self.assertEqual(updated.candidates[2], before.candidates[2])
        self.assertEqual([c['가격'] for c in updated.candidates], [110, 125, 130])
        self.assertEqual(updated.excluded, [0])
        self.assertIn(target, updated.rejected_ids)
        self.assertEqual(self.client.post(f'/jobs/{job_id}/items/0/research', json={
            'revision': 2, 'replace_item_id': '30000000'}).status_code, 202)
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=ReplacementProvider())
        self.assertEqual(self.store.item(job_id, 0).candidates, updated.candidates)
        self.assertIn('기존 후보를 유지', self.store.read(job_id).message)
        self.assertEqual(self.client.post(f'/jobs/{job_id}/workbook', json={'revision': 3}).status_code, 202)
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=Provider())
        book = load_workbook(io.BytesIO(self.client.get(f'/jobs/{job_id}/workbook').content))
        self.addCleanup(book.close)
        self.assertEqual([book.worksheets[0].cell(9, col).value for col in (5, 6)], ['125', '130'])

    def test_research_rejects_invalid_requests_and_preserves_results_on_failure(self):
        job_id = self.create()
        self.collect(job_id)
        before = self.store.item(job_id, 0)
        path = f'/jobs/{job_id}/items/0/research'
        target = before.candidates[0]['식별번호']
        for body, status in [
            ({'revision': 99, 'replace_item_id': target}, 409),
            ({'revision': 0, 'replace_item_id': '99999999'}, 422),
            ({'revision': 0, 'replace_item_id': target, 'conditions': {'arbitrary': True}}, 422),
            ({'revision': 0, 'replace_item_id': target, 'conditions': {'material': ', /'}}, 422),
        ]:
            self.assertEqual(self.client.post(path, json=body).status_code, status)
        response = self.client.post(path, json={'revision': 0, 'replace_item_id': target})
        self.assertEqual(response.status_code, 202, response.text)
        self.assertEqual(self.client.post(path, json={'revision': 1, 'replace_item_id': target}).status_code, 409)
        class OfflineProvider(Provider):
            def search(self, *args):
                raise requests.ConnectionError('offline')
        with self.assertRaises(requests.ConnectionError):
            execute(self.root, job_id, response.json()['operation_id'], provider=OfflineProvider())
        self.assertEqual(self.store.item(job_id, 0), before)
        self.assertEqual(self.client.post(f'/jobs/{job_id}/retry').status_code, 202)
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=Provider())
        self.assertEqual(self.store.item(job_id, 0).candidates, before.candidates)
        self.assertEqual(self.store.read(job_id).status, 'review')

    def test_research_requires_requested_material_grade_without_relaxing_price_or_dimensions(self):
        job_id = self.create()
        raw = product('10000000', '기준회사', 100)['_raw']
        raw['pdctAtrbCdDtlNm'] = '재질:목재/SUS304,용도:방향안내'
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=Provider(),
                collect=lambda request: {'기준회사': [extract(raw)]})
        before = self.store.item(job_id, 0)
        response = self.client.post(f'/jobs/{job_id}/items/0/research', json={
            'revision': 0, 'replace_item_id': before.candidates[0]['식별번호'],
            'conditions': {'material': 'STS304'}})
        self.assertEqual(response.status_code, 202, response.text)
        class MaterialProvider(Provider):
            def search(self, *args):
                rows = []
                for index, (material, price, dimensions) in enumerate((
                    ('목재', 101, '600×150'), ('SUS316', 102, '600×150'),
                    ('스테인리스', 103, '600×150'), ('SUS304', 99, '600×150'),
                    ('SUS304', 104, '900×150'), ('SUS304/알루미늄', 125, '600×150'),
                )):
                    row_ = product(str(30000000+index), f'대체{index}', price)['_raw']
                    row_['pdctAtrbCdDtlNm'] = f'재질:{material},용도:방향안내'
                    row_['itemIdnfNm'] = f'안내판, MODEL, {dimensions}mm'
                    rows.append(row_)
                return CandidatePage(tuple(rows), len(rows), False, None)
        execute(self.root, job_id, response.json()['operation_id'], provider=MaterialProvider())
        self.assertEqual([c['식별번호'] for c in self.store.item(job_id, 0).candidates],
                         [before.candidates[1]['식별번호'], '30000005', before.candidates[2]['식별번호']])
        operation = self.store.read(job_id).operation_id
        saved = self.store.item(job_id, 0)
        self.store.update(job_id, lambda job: job.model_copy(update={'status': 'running'}))
        class UnexpectedProvider(Provider):
            def search(self, *args):
                raise AssertionError('committed research must not run again')
        execute(self.root, job_id, operation, provider=UnexpectedProvider())
        self.assertEqual(self.store.item(job_id, 0), saved)

    def test_cancel_during_research_does_not_replace_candidate(self):
        job_id = self.create()
        self.collect(job_id)
        before = self.store.item(job_id, 0)
        response = self.client.post(f'/jobs/{job_id}/items/0/research', json={
            'revision': 0, 'replace_item_id': before.candidates[0]['식별번호']})
        self.assertEqual(response.status_code, 202, response.text)
        client = self.client
        class CancellingProvider(Provider):
            def search(self, *args):
                client.post(f'/jobs/{job_id}/cancel')
                return CandidatePage((product('30000000', '새 업체', 125)['_raw'],), 1, False, None)
        execute(self.root, job_id, response.json()['operation_id'], provider=CancellingProvider())
        self.assertEqual(self.store.item(job_id, 0), before)
        self.assertEqual(self.store.read(job_id).status, 'cancelled')

    def test_frontend_assets_revalidate_after_deployment(self):
        for path in ('/', '/app.js', '/views.js'):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.headers['cache-control'], 'no-cache')
                cached = self.client.get(path, headers={'If-None-Match': response.headers['etag']})
                self.assertEqual(cached.status_code, 304)
                self.assertEqual(cached.headers['cache-control'], 'no-cache')
        self.assertEqual(self.client.get('/jobs').headers['cache-control'], 'no-store')

    def test_optional_direction_reaches_worker_comparison(self):
        class DirectionProvider(Provider):
            def search(self, keyword, lower='', upper='', limit=500):
                raw = product('20000000', '경쟁사', 110)['_raw']
                raw['pdctAtrbCdDtlNm'] = '재질:목재,용도:2방향안내,지지대수:2개'
                return CandidatePage((raw,), 1, False, None)

        for strict, expected in ((False, 1), (True, 0)):
            with self.subTest(same_direction=strict):
                response = self.client.post('/jobs', json={
                    'company': '기준회사', 'same_direction': strict})
                self.assertEqual(response.status_code, 202)
                job = self.store.read(response.json()['id'])
                self.assertEqual(job.request.limit, 0)
                reference = product('10000000', '기준회사', 100)
                reference['_raw']['pdctAtrbCdDtlNm'] = '재질:목재,용도:1방향안내,지지대수:2개'
                reference = extract(reference['_raw'])
                execute(self.root, job.id, job.operation_id, provider=DirectionProvider(),
                        collect=lambda request, reference=reference: {'기준회사': [reference]})
                self.assertEqual(len(self.store.items(job.id)[0].candidates), expected)

    def test_recording_mode_exports_three_columns_without_inventing_support_counts(self):
        class RecordingProvider(Provider):
            def search(self, keyword, lower='', upper='', limit=500):
                return CandidatePage(tuple(product(str(20000000+i), f'경쟁사{i}', price, count)['_raw']
                    for i, (price, count) in enumerate(((130, ''), (110, '1개'), (120, '2개')))), 3, False, None)

            def enrich(self, raw, *, required=()):
                if '지지대수' in required:
                    raise AssertionError('disabled support check requested enrichment')
                return raw

        job_id = self.create()
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=RecordingProvider(),
                collect=lambda request: {'기준회사': [product('10000000', '기준회사', 100, '')]})
        item = self.store.items(job_id)[0]
        self.assertIsNone(item.reference['지지대수'])
        self.assertNotIn('비교보류사유', item.reference)
        self.assertEqual([c['가격'] for c in item.candidates], [110, 120, 130])
        self.assertTrue(all('지지대 검사 생략' in c['매칭근거'] for c in item.candidates))
        self.assertIsNone(item.candidates[-1]['지지대수'])
        response = self.client.post(f'/jobs/{job_id}/workbook', json={'revision': 0})
        self.assertEqual(response.status_code, 202)
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=RecordingProvider())
        download = self.client.get(f'/jobs/{job_id}/workbook')
        book = load_workbook(io.BytesIO(download.content))
        self.addCleanup(book.close)
        sheet = book.worksheets[0]
        self.assertEqual([sheet.cell(9, col).value for col in (4, 5, 6, 7)], ['100', '110', '120', '130'])

    def test_review_labels_survive_storage_and_excel_export(self):
        class ReviewProvider(Provider):
            def search(self, *args):
                rows = []
                for i, (price, face) in enumerate(((130, '양면'), (110, ''), (120, '단면'))):
                    raw = product(str(20000000+i), f'경쟁사{i}', price)['_raw']
                    raw['itemIdnfNm'] += ', ' + face
                    raw['dtlsPrnmNo'] = '5512171801'
                    rows.append(raw)
                return CandidatePage(tuple(rows), 3, False, None)
        reference = product('10000000', '기준회사', 100)['_raw']
        reference['itemIdnfNm'] += ', 양면'
        reference['dtlsPrnmNo'] = '5512171801'
        job_id = self.create()
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=ReviewProvider(),
                collect=lambda request: {'기준회사': [extract(reference)]})
        item = self.client.get(f'/jobs/{job_id}/items').json()[0]
        self.assertEqual([c['비교구분'] for c in item['candidates']], ['확인 필요', '대안 후보', '조건 충족'])
        self.assertEqual(item['reference']['후보진단']['검사후보'], 3)
        self.client.post(f'/jobs/{job_id}/workbook', json={'revision': 0})
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=ReviewProvider())
        book = load_workbook(io.BytesIO(self.client.get(f'/jobs/{job_id}/workbook').content))
        self.addCleanup(book.close)
        sheet = book.worksheets[0]
        self.assertEqual([sheet.cell(3, col).value for col in (5, 6, 7)],
                         ['경쟁사1 [확인 필요]', '경쟁사2 [대안 후보]', '경쟁사0'])
        self.assertIn('표시면수 미확인', sheet['E3'].comment.text)
        self.assertIn('표시면수 차이: 2 → 1', sheet['F3'].comment.text)
        self.assertEqual([sheet.cell(9, col).value for col in (5, 6, 7)], ['110', '120', '130'])

    def test_review_persists_and_export_uses_reviewed_selection_without_search(self):
        job_id = self.create()
        self.collect(job_id)
        item = self.client.get(f'/jobs/{job_id}/items').json()[0]
        self.assertEqual([c['가격'] for c in item['candidates']], ['110', '120', '130'])
        revision = self.store.read(job_id).revision
        response = self.client.post(f'/jobs/{job_id}/items/0/selection',
                                    json={'revision': revision, 'excluded': [1]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(JobStore(self.root).items(job_id)[0].excluded, [1])
        response = self.client.post(f'/jobs/{job_id}/workbook', json={'revision': revision+1})
        self.assertEqual(response.status_code, 202, response.text)
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=Provider())
        download = self.client.get(f'/jobs/{job_id}/workbook')
        self.assertEqual(download.status_code, 200)
        self.assertEqual(self.store.read(job_id).completed, self.store.read(job_id).total)
        workbook = load_workbook(io.BytesIO(download.content))
        self.addCleanup(workbook.close)
        values = [c.value for row_ in workbook.worksheets[0] for c in row_]
        self.assertIn('110', values)
        self.assertIn('130', values)
        self.assertNotIn('120', values)
        response = self.client.post(f'/jobs/{job_id}/items/0/selection',
                                    json={'revision': revision+1, 'excluded': []})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.get(f'/jobs/{job_id}/workbook').status_code, 409)

    def test_collection_retry_archives_previous_items_when_source_shrinks(self):
        job_id = self.create()
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=Provider(),
                collect=lambda request: {'기준회사': [product('10000000', '기준회사', 100),
                                                     product('10000001', '기준회사', 100)]})
        self.store.update(job_id, lambda job: job.model_copy(update={'status': 'failed'}))
        self.assertEqual(self.client.post(f'/jobs/{job_id}/retry').status_code, 202)
        operation = self.store.read(job_id).operation_id
        self.collect(job_id)
        self.assertEqual(len(self.store.items(job_id)), 1)
        self.assertEqual(self.store.read(job_id).total, 1)
        self.assertEqual(len(list((self.root/job_id/f'items-before-{operation}').glob('*.json'))), 2)

    @patch('core.comparison.CHECK_SUPPORTS', True)
    def test_human_support_evidence_recompares_and_keeps_source_unknown(self):
        job_id = self.create()
        self.collect(job_id, supports='')
        self.assertEqual(self.store.items(job_id)[0].candidates, [])
        response = self.client.post(f'/jobs/{job_id}/items/0/supports',
                                    json={'revision': 0, 'count': 2, 'note': '규격서 3쪽 확인'})
        self.assertEqual(response.status_code, 202, response.text)
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=Provider())
        item = self.store.items(job_id)[0]
        self.assertEqual(len(item.candidates), 3)
        self.assertEqual(item.reference['지지대수'], 2)
        self.assertIsNone(extract(item.reference['_raw'])['지지대수'])
        assert item.support_review is not None
        self.assertEqual(item.support_review.count, 2)

    def test_invalid_stale_and_cross_origin_writes_are_rejected(self):
        self.assertEqual(self.client.post('/jobs', json={'company': ' '}).status_code, 422)
        self.assertEqual(self.client.post('/jobs', json={'company': 'a'},
                         headers={'Origin': 'https://untrusted.example'}).status_code, 403)
        job_id = self.create()
        self.assertEqual(self.client.post(f'/jobs/{job_id}/workbook', json={'revision': 0}).status_code, 409)
        self.collect(job_id)
        self.assertEqual(self.client.post(f'/jobs/{job_id}/items/0/selection',
                         json={'revision': 99, 'excluded': []}).status_code, 409)
        self.assertEqual(self.store.items(job_id)[0].reference['가격'], Decimal(100))

    def test_cancelled_queued_job_does_not_fetch(self):
        job_id = self.create()
        operation = self.store.read(job_id).operation_id
        self.assertEqual(self.client.post(f'/jobs/{job_id}/cancel').status_code, 200)
        def unexpected(request):
            self.fail('cancelled job fetched data')
        execute(self.root, job_id, operation, collect=unexpected)
        self.assertEqual(self.store.read(job_id).status, 'cancelled')

    def test_duplicate_delivery_does_not_repeat_completed_work(self):
        job_id = self.create()
        self.collect(job_id)
        def unexpected(request):
            self.fail('completed job fetched data again')
        execute(self.root, job_id, self.store.read(job_id).operation_id, collect=unexpected)
        self.assertEqual(len(self.store.items(job_id)), 1)

    def test_source_conflict_is_rejected_without_rewriting_evidence(self):
        job_id = self.create()
        self.collect(job_id)
        response = self.client.post(f'/jobs/{job_id}/items/0/supports',
                                   json={'revision': 0, 'count': 1, 'note': '다른 수량 입력'})
        self.assertEqual(response.status_code, 409)
        self.assertIsNone(self.store.items(job_id)[0].support_review)

    def test_source_error_is_visible_and_retry_gets_new_operation(self):
        job_id = self.create()
        operation = self.store.read(job_id).operation_id
        def unavailable(request):
            raise requests.ConnectionError('fixture offline')
        with self.assertRaises(requests.ConnectionError):
            execute(self.root, job_id, operation, collect=unavailable)
        self.assertEqual(self.client.get(f'/jobs/{job_id}').json()['status'], 'failed')
        response = self.client.post(f'/jobs/{job_id}/retry')
        self.assertEqual(response.status_code, 202)
        self.assertNotEqual(response.json()['operation_id'], operation)

    def test_broker_outage_keeps_job_for_retry(self):
        def offline(job_id, operation_id):
            raise requests.ConnectionError('fixture broker offline')
        client = TestClient(create_app(self.root, submit=offline))
        response = client.post('/jobs', json={'company': '기준회사'})
        self.assertEqual(response.status_code, 503)
        job_id = response.json()['detail']['job_id']
        self.assertEqual(self.store.read(job_id).status, 'queued')
        self.assertEqual(self.client.post(f'/jobs/{job_id}/retry').status_code, 202)

    def test_invalid_upload_and_unknown_job_are_rejected(self):
        self.assertEqual(self.client.post('/jobs/upload', content=b'not an xls').status_code, 422)
        self.assertEqual(self.client.get('/jobs/00000000-0000-0000-0000-000000000000/items').status_code, 404)
        self.assertEqual(self.store.recent(), [])

    def test_decimal_price_and_snapshot_survive_file_reload(self):
        job_id = self.create()
        original = product('10000000', '기준회사', 100)
        original['가격'] = Decimal('123456789012345678.1234')
        original['_raw'] = {**original['_raw'], 'ctrtUprc': '123456789012345678.1234'}
        execute(self.root, job_id, self.store.read(job_id).operation_id, provider=Provider(),
                collect=lambda request: {'기준회사': [original]})
        self.assertEqual(JobStore(self.root).items(job_id)[0].reference['가격'], Decimal('123456789012345678.1234'))

    def test_prefect_crash_changes_stuck_job_to_retryable_failure(self):
        job_id = self.create()
        self.store.update(job_id, lambda job: job.model_copy(update={'flow_run_id': 'test-flow'}))
        with patch('src.web.dispatch.prefect_client') as client:
            client.return_value.__enter__.return_value.get.return_value.json.return_value = {'state_type': 'CRASHED'}
            self.assertEqual(self.client.get(f'/jobs/{job_id}').json()['status'], 'failed')


if __name__ == '__main__':
    unittest.main()
