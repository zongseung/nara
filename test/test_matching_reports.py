import json
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from core.matching.contracts import NORMALIZER_VERSION, Purpose, ReleaseManifest
from core.matching.policy import load_policy
from src.matching.releases import Release, write_release
from test.test_matching_policy import policy_data
from test.updater_fixtures import FixtureSource, row
from src.updater.collector import CollectionPolicy, refresh
from src.updater.store import Store


def prepare(root):
    source = Store(root / 'source')
    manifest = refresh(source, FixtureSource([row('A', spec='20x30mm')]), policy=CollectionPolicy(request_interval=0), environment='fixture')
    manifest_path = root / 'manifest.json'
    manifest_path.write_text(manifest.model_dump_json())
    policies = root / 'policies'
    policies.mkdir()
    payload = policy_data(category_ids=['123'], state='active')
    (policies / 'policy.json').write_text(json.dumps(payload))
    policy = load_policy(payload)
    release = Release(manifest=ReleaseManifest('fixture', Purpose.FUNCTIONAL_SUBSTITUTE, '123', policy.content_hash, 'fixture-dataset', manifest.source_normalizer_version, NORMALIZER_VERSION, 'fixture'))
    write_release(release, root / 'release.json')
    book = Workbook()
    book.active.title = '회사'
    book.active.append(['물품식별번호', '한글품명', '희망가(부가세포함)'])
    book.active.append(['00001', '품목 20x30mm', 100])
    book.save(root / 'input.xlsx')
    return ['--purpose', 'functional_substitute', '--category', '123', '--catalog-manifest', str(manifest_path), '--source-root', str(source.root), '--policy-dir', str(policies), '--release', str(root / 'release.json'), '--ledger', str(root / 'matching.sqlite'), '--index-dir', str(root / 'index')]


class ReportIntegration(unittest.TestCase):
    def test_real_subprocess_workbook_and_ledger(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            options = prepare(root)
            command = [sys.executable, 'compare_file.py', str(root / 'input.xlsx'), '--mode', 'shadow', *options, '--out', str(root / 'report.xlsx')]
            result = subprocess.run(command, check=False, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            workbook = load_workbook(root / 'report.xlsx')
            self.addCleanup(workbook.close)
            self.assertTrue({'판정요약', '검토필요', '판정근거', '데이터버전', '업체_회사'} <= set(workbook.sheetnames))
            self.assertEqual(workbook['판정요약']['C2'].value, 'review_required')
            self.assertEqual(workbook['업체_회사']['J2'].value, 100)
            metadata = dict(workbook['데이터버전'].values)
            self.assertFalse(json.loads(metadata['source'])['source_snapshot_verified'])
            with sqlite3.connect(root / 'matching.sqlite') as db:
                self.assertEqual(db.execute('SELECT count(*) FROM decisions').fetchone()[0], 1)
            result = subprocess.run([sys.executable, 'compare_file.py', str(root / 'input.xlsx'), '--mode', 'auto', *options, '--out', str(root / 'auto.xlsx')], check=False, capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse((root / 'auto.xlsx').exists())

    def test_text_safety_and_row_splitting(self):
        from dataclasses import replace

        from src.matching.engine import run_match
        from src.matching.reporting import write_report
        from src.matching.store import MatchingStore
        from test.test_matching_engine import NOW, pipeline
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with pipeline(root) as (item, policy, release, index), MatchingStore(root / 'ledger.sqlite') as ledger:
                item = replace(item, name_raw='=HYPERLINK("https://example.invalid")', sheet='x' * 40)
                result = run_match(item, policy, index, release=release, now=NOW)
                job = ledger.record_job(item, result, release)
                with patch('src.matching.reporting.MAX_ROWS', 3):
                    write_report(root / 'report.xlsx', [(item, result, job)] * 4, release=release, manifest=index.catalog.manifest, observed_at=NOW, mode='shadow', connection=ledger.connection)
            workbook = load_workbook(root / 'report.xlsx')
            self.addCleanup(workbook.close)
            self.assertTrue(all(len(sheet.title) <= 31 and sheet.max_row <= 3 for sheet in workbook))
            company = next(sheet for sheet in workbook if sheet.title.startswith('업체_'))
            self.assertEqual(company['B2'].data_type, 's')
            self.assertTrue(company['B2'].value.startswith('='))

    def test_auto_selected_and_unselected_are_distinct_without_exception_rows(self):
        from dataclasses import replace

        from core.matching.contracts import DecisionStatus
        from src.matching.engine import run_match
        from src.matching.reporting import write_report
        from src.matching.store import MatchingStore
        from test.test_matching_engine import NOW, pipeline
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with pipeline(root) as (item, policy, release, index), MatchingStore(root/'ledger.sqlite') as ledger:
                original = run_match(item, policy, index, release=release, now=NOW)
                check = original.candidate_checks[0]
                extra = replace(check, candidate=replace(check.candidate, listing_key='unselected'))
                result = replace(original, decision=replace(original.decision, status=DecisionStatus.AUTO_CONFIRMED, candidates=(check.candidate.listing_key,)), candidate_checks=(check, extra))
                job = ledger.record_job(item, result, release)
                write_report(root/'report.xlsx', [(item, result, job)], release=release, manifest=index.catalog.manifest, observed_at=NOW, mode='auto', connection=ledger.connection)
            workbook = load_workbook(root/'report.xlsx')
            self.addCleanup(workbook.close)
            company = next(sheet for sheet in workbook if sheet.title.startswith('업체_'))
            self.assertEqual(company['T2'].value, '자동선택')
            self.assertEqual(company['T3'].value, '미선택')
            self.assertFalse(any(any(value is not None for value in row) for row in list(workbook['검토필요'].values)[1:]))

    def test_auto_mode_review_selection_is_never_green_automatic(self):
        from dataclasses import replace

        from src.matching.engine import run_match
        from src.matching.reporting import write_report
        from src.matching.store import MatchingStore
        from test.test_matching_engine import NOW, pipeline
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with pipeline(root) as (item, policy, release, index), MatchingStore(root/'ledger.sqlite') as ledger:
                original = run_match(item, policy, index, release=release, now=NOW)
                result = replace(original, decision=replace(original.decision, candidates=(original.candidate_checks[0].candidate.listing_key,)))
                job = ledger.record_job(item, result, release)
                write_report(root/'report.xlsx', [(item, result, job)], release=release, manifest=index.catalog.manifest, observed_at=NOW, mode='auto', connection=ledger.connection)
            workbook = load_workbook(root/'report.xlsx')
            self.addCleanup(workbook.close)
            company = next(sheet for sheet in workbook if sheet.title.startswith('업체_'))
            self.assertEqual(workbook['판정요약']['C2'].value, 'review_required')
            self.assertEqual(company['T2'].value, '검토후보')
            self.assertNotEqual(company['T2'].fill.fgColor.rgb, '00C6EFCE')
