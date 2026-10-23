import json
import os
from pathlib import Path
import subprocess
import sys
import sqlite3
import tempfile
import unittest
from openpyxl import load_workbook
from src.updater.collector import refresh, CollectionPolicy
from src.updater.exports import export_xlsx
from src.updater.locks import file_lock
from src.updater.reader import open_published_snapshot, Manifest
from src.updater.store import Store, StoreError
from test.updater_fixtures import FixtureSource, row


class ReaderTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = Store(Path(self.temp.name))

    def publish(self, price='100'):
        return refresh(self.store, FixtureSource([row('=1+1', price)]), policy=CollectionPolicy(request_interval=0), environment='fixture')

    def test_pin_survives_publication_and_stale_manifest_rejected(self):
        first = self.publish()
        with open_published_snapshot(first, root=self.store.root) as reader:
            second = self.publish('200')
            self.assertEqual(list(reader.current())[0]['contract_price_decimal'], '100')
            with self.assertRaises(BlockingIOError):
                with file_lock(self.store.path(first.reader_locator.pin), blocking=False):
                    self.fail('reader pin was lost')
            with self.assertRaises(StoreError):
                with open_published_snapshot(first, root=self.store.root):
                    self.fail('stale manifest accepted')
        with open_published_snapshot(second, root=self.store.root) as reader:
            self.assertEqual(list(reader.current())[0]['contract_price_decimal'], '200')
            with self.assertRaises(sqlite3.OperationalError) as error:
                reader.connection.execute('DELETE FROM items_current')
            self.assertIn('readonly', str(error.exception))

    def test_hash_and_symlink_rejected(self):
        manifest = self.publish()
        path = self.store.path(manifest.snapshot_path)
        with path.open('ab') as handle:
            handle.write(b'corruption')
        with self.assertRaises(StoreError):
            with open_published_snapshot(manifest, root=self.store.root):
                self.fail('hash mismatch accepted')
        moved = path.with_suffix('.moved')
        path.rename(moved)
        path.symlink_to(moved)
        with self.assertRaises(StoreError):
            with open_published_snapshot(manifest, root=self.store.root):
                self.fail('symlink accepted')

    def test_xlsx_formula_text_and_kst(self):
        manifest = self.publish()
        destination = self.store.path('export.xlsx')
        with open_published_snapshot(manifest, root=self.store.root) as snapshot:
            export_xlsx(snapshot, destination, sheet_rows=2)
        book = load_workbook(destination, read_only=True, data_only=False)
        self.addCleanup(book.close)
        values = list(book.active.values)
        self.assertEqual(values[1][0], 'shopping:listing:ctrtItemMngNo:=1+1')
        self.assertIn('[Asia/Seoul]', values[1][12])
        self.assertEqual(values[1][6], '100')
        self.assertEqual(book.active['A2'].data_type, 's')

    def test_cli_real_subprocess_current_manifest_history_export(self):
        manifest = self.publish()
        base = [sys.executable, '-m', 'src.updater.cli', '--data-dir', str(self.store.root)]
        for command in (['status','--json'], ['current'], ['history','shopping:listing:ctrtItemMngNo:=1+1','--price-only'],
                        ['catalog-manifest','--out',str(self.store.path('manifest.json'))], ['export','--out',str(self.store.path('cli.xlsx'))]):
            result = subprocess.run(base + command, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(Manifest.model_validate_json(self.store.path('manifest.json').read_text()), manifest)
        self.assertGreater(self.store.path('cli.xlsx').stat().st_size, 0)

    def test_cli_unpublished_commands_explain_missing_data(self):
        base = [sys.executable, '-m', 'src.updater.cli', '--data-dir', str(self.store.root)]
        for command in (['current'], ['history', 'missing-listing'],
                        ['catalog-manifest', '--out', str(self.store.path('manifest.json'))],
                        ['export', '--out', str(self.store.path('catalog.xlsx'))]):
            with self.subTest(command=command[0]):
                result = subprocess.run(base + command, capture_output=True, text=True)
                self.assertEqual(result.returncode, 2, result.stderr)
                self.assertIn('no_publication', result.stderr)
                self.assertIn('status --json', result.stderr)
                self.assertNotIn('Traceback', result.stderr)
                self.assertEqual(result.stdout, '')
        self.assertFalse(self.store.path('manifest.json').exists())
        self.assertFalse(self.store.path('catalog.xlsx').exists())
        status = subprocess.run(base + ['status', '--json'], capture_output=True, text=True)
        self.assertEqual(status.returncode, 0, status.stderr)
        self.assertEqual(json.loads(status.stdout)['freshness']['status'], 'unpublished')

    def test_cli_refresh_disabled_does_not_enqueue(self):
        result = subprocess.run([sys.executable, '-m', 'src.updater.cli', '--data-dir',
                                 str(self.store.root), 'refresh'], capture_output=True, text=True,
                                env=os.environ | {'NARA_SOURCE_ENABLED': 'false'})
        self.assertEqual(result.returncode, 2, result.stderr)
        self.assertIn('production_source_not_enabled', result.stderr)
        self.assertNotIn('Traceback', result.stderr)
        with self.store.control() as conn:
            self.assertEqual(conn.execute('SELECT count(*) FROM runs').fetchone()[0], 0)

    def test_cli_status_reports_published_collection_progress(self):
        self.publish()
        result = subprocess.run([sys.executable, '-m', 'src.updater.cli', '--data-dir',
                                 str(self.store.root), 'status', '--json'], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(report['collection_progress']['expected_items'], 1)
        self.assertEqual(report['collection_progress']['collected_items'], 1)
        self.assertEqual(report['collection_progress']['completed_pages'], 1)
        self.assertEqual(report['collection_progress']['expected_pages'], 1)

    def test_export_includes_history_and_contract_period(self):
        self.publish('100')
        manifest = self.publish('120')
        destination = self.store.path('history.xlsx')
        with open_published_snapshot(manifest, root=self.store.root) as snapshot:
            export_xlsx(snapshot, destination, sheet_rows=2)
        book = load_workbook(destination,read_only=True)
        self.addCleanup(book.close)
        self.assertTrue(any(name.startswith('가격이력') for name in book.sheetnames))
        self.assertTrue(any(name.startswith('필드변경') for name in book.sheetnames))
        self.assertIn('contract_start',next(book['상품1'].values))
